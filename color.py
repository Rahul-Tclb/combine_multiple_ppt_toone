from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_slide_with_table(prs, title, table_data):
    """
    Creates a slide with a table containing the given data.

    Args:
        prs: The presentation object.
        title: The title of the slide.
        table_data: A list of lists representing the table data.

    Returns:
        The created slide.
    """
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    shapes = slide.shapes

    # Title
    title_left = Inches(1)
    title_top = Inches(1)
    title_width = Inches(10)
    title_height = Inches(0.5)
    txBox = shapes.add_textbox(title_left, title_top, title_width, title_height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True

    # Table
    table_left = Inches(1)
    table_top = Inches(2)
    table_width = Inches(10)
    table_height = Inches(1)  # Adjusted for better appearance
    rows, cols = len(table_data), len(table_data[0])
    table = shapes.add_table(rows, cols, table_left, table_top, table_width, table_height).table

    # Set table header style
    for j, header_text in enumerate(table_data[0]):
        cell = table.cell(0, j)
        if header_text == "TPA Involved (Registered, Un-Registered)":
            # Create formatted text for this column header
            cell.text = ""
            text_frame = cell.text_frame
            p = text_frame.add_paragraph()
            p.font.bold = True
            p.font.size = Pt(12)

            # Add "TPA Involved"
            run = p.add_run()
            run.text = "TPA Involved ("

            # Add "Registered" in green
            run = p.add_run()
            run.text = "Registered"
            run.font.color.rgb = RGBColor(0, 255, 0)

            # Add ", " and "Un-Registered" in red
            run = p.add_run()
            run.text = ", "

            run = p.add_run()
            run.text = "Un-Registered"
            run.font.color.rgb = RGBColor(255, 0, 0)

            # Add closing parenthesis
            run = p.add_run()
            run.text = ")"
        else:
            # Standard header formatting
            cell.text = header_text
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(12)

    # Set table cell styles
    for i, row in enumerate(table_data[1:], start=1):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(cell_text)
            cell.text_frame.paragraphs[0].font.size = Pt(10)

    return slide

# Example usage:
prs = Presentation()
table_data = [
    ["No of cases", "Program", "Total NCA", "TPA Involved (Registered, Un-Registered)", "HIR License", "Warning Letter"],
    ["123", "Program A", "100", "TPA Involved", "Yes", "No"],
    ["456", "Program B", "200", "(Registered, UnRegistered)", "No", "Yes"]
]

create_slide_with_table(prs, "My Slide Title", table_data)

# Save the presentation
prs.save('output.pptx')
